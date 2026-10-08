import io
from dataclasses import dataclass
from typing import Literal
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate

from app.schemas.profile import ProfileData

LineStyle = Literal["name", "contact", "heading", "subheading", "body", "bullet"]


@dataclass(frozen=True)
class ResumeLine:
    """A single text line with a semantic role, shared by both renderers."""

    text: str
    style: LineStyle


def build_lines(profile: ProfileData) -> list[ResumeLine]:
    lines: list[ResumeLine] = [ResumeLine(profile.full_name or "Resume", "name")]

    contact = " · ".join(
        str(part) for part in (profile.email, profile.phone, profile.location) if part
    )
    if contact:
        lines.append(ResumeLine(contact, "contact"))

    if profile.summary:
        lines += _section("Summary", [ResumeLine(profile.summary, "body")])
    if profile.work_history:
        lines += _section("Experience", _experience_lines(profile))
    if profile.education:
        lines += _section("Education", _education_lines(profile))
    if profile.skills:
        lines += _section("Skills", [ResumeLine(", ".join(profile.skills), "body")])
    if profile.projects:
        lines += _section("Projects", _project_lines(profile))
    if profile.certifications:
        lines += _section(
            "Certifications", [ResumeLine(item, "bullet") for item in profile.certifications]
        )
    if profile.achievements:
        lines += _section(
            "Achievements", [ResumeLine(item, "bullet") for item in profile.achievements]
        )
    return lines


def _section(heading: str, body: list[ResumeLine]) -> list[ResumeLine]:
    return [ResumeLine(heading, "heading"), *body]


def _experience_lines(profile: ProfileData) -> list[ResumeLine]:
    lines: list[ResumeLine] = []
    for job in profile.work_history:
        title = " — ".join(part for part in (job.title, job.company) if part)
        if title:
            lines.append(ResumeLine(title, "subheading"))
        dates = _range(job.start_date, job.end_date)
        if dates:
            lines.append(ResumeLine(dates, "body"))
        if job.description:
            lines.append(ResumeLine(job.description, "body"))
    return lines


def _education_lines(profile: ProfileData) -> list[ResumeLine]:
    lines: list[ResumeLine] = []
    for education in profile.education:
        title = " — ".join(part for part in (education.degree, education.institution) if part)
        lines.append(ResumeLine(title or education.institution, "subheading"))
        details = [education.field_of_study] if education.field_of_study else []
        years = _range(
            str(education.start_year) if education.start_year else None,
            str(education.end_year) if education.end_year else None,
        )
        if years:
            details.append(years)
        if details:
            lines.append(ResumeLine(" · ".join(details), "body"))
    return lines


def _project_lines(profile: ProfileData) -> list[ResumeLine]:
    lines: list[ResumeLine] = []
    for project in profile.projects:
        lines.append(ResumeLine(project.name, "subheading"))
        if project.description:
            lines.append(ResumeLine(project.description, "body"))
        if project.technologies:
            lines.append(ResumeLine(", ".join(project.technologies), "body"))
        if project.url:
            lines.append(ResumeLine(project.url, "body"))
    return lines


def _range(start: str | None, end: str | None) -> str | None:
    if not start and not end:
        return None
    return f"{start or ''} – {end or 'Present'}".strip()


def render_pdf(profile: ProfileData) -> bytes:
    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=LETTER,
        title=profile.full_name or "Resume",
        author=profile.full_name or "",
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.6 * inch,
    )
    story: list[Paragraph] = []
    for line in build_lines(profile):
        if line.style == "bullet":
            story.append(Paragraph(escape(line.text), _PDF_STYLES["bullet"], bulletText="•"))
        else:
            story.append(Paragraph(escape(line.text), _PDF_STYLES[line.style]))
    document.build(story)
    return buffer.getvalue()


def render_docx(profile: ProfileData) -> bytes:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    for line in build_lines(profile):
        paragraph = document.add_paragraph()
        run = paragraph.add_run(f"• {line.text}" if line.style == "bullet" else line.text)
        if line.style == "name":
            run.bold = True
            run.font.size = Pt(18)
        elif line.style == "contact":
            run.font.size = Pt(10)
        elif line.style == "heading":
            run.bold = True
            run.font.size = Pt(12)
        elif line.style == "subheading":
            run.bold = True
        elif line.style == "bullet":
            paragraph.paragraph_format.left_indent = Inches(0.25)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


_PDF_STYLES: dict[str, ParagraphStyle] = {
    "name": ParagraphStyle("name", fontName="Helvetica-Bold", fontSize=16, leading=20),
    "contact": ParagraphStyle(
        "contact", fontName="Helvetica", fontSize=10, leading=13, spaceAfter=8
    ),
    "heading": ParagraphStyle(
        "heading", fontName="Helvetica-Bold", fontSize=11, leading=14, spaceBefore=10, spaceAfter=4
    ),
    "subheading": ParagraphStyle(
        "subheading", fontName="Helvetica-Bold", fontSize=10, leading=13, spaceBefore=4
    ),
    "body": ParagraphStyle("body", fontName="Helvetica", fontSize=10, leading=13, spaceAfter=2),
    "bullet": ParagraphStyle(
        "bullet",
        fontName="Helvetica",
        fontSize=10,
        leading=13,
        leftIndent=12,
        bulletIndent=0,
        spaceAfter=1,
    ),
}
