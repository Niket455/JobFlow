import io

import pdfplumber
from docx import Document
from fastapi.testclient import TestClient

from app.core.media_types import DOCX_CONTENT_TYPE, PDF_CONTENT_TYPE
from app.schemas.profile import ProfileData
from app.services.resume_renderer import render_docx, render_pdf

USER = {"email": "export@example.com", "password": "supersecret1"}
PROFILE = ProfileData(
    full_name="Jane Doe",
    email="jane@example.com",
    phone="+91 90000 00000",
    location="Bengaluru",
    summary="Data engineer",
    skills=["Python", "SQL", "PySpark"],
)


def _pdf_text(data: bytes) -> str:
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


def _docx_text(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def _docx_bytes(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _headers(client: TestClient) -> dict[str, str]:
    client.post("/api/v1/auth/register", json=USER)
    tokens = client.post("/api/v1/auth/login", json=USER).json()
    return {"Authorization": f"Bearer {tokens['access_token']}"}


def _upload(client: TestClient, headers: dict[str, str]) -> str:
    files = {"file": ("resume.docx", _docx_bytes("Jane Doe Python SQL"), DOCX_CONTENT_TYPE)}
    response = client.post("/api/v1/resumes", files=files, headers=headers)
    assert response.status_code == 201
    return str(response.json()["id"])


def test_render_pdf_extracts_text() -> None:
    data = render_pdf(PROFILE)
    assert data.startswith(b"%PDF")
    text = _pdf_text(data)
    assert "Jane Doe" in text
    assert "Python" in text


def test_render_docx_extracts_text() -> None:
    data = render_docx(PROFILE)
    text = _docx_text(data)
    assert "Jane Doe" in text
    assert "Python" in text


def test_export_requires_auth(client: TestClient) -> None:
    response = client.get("/api/v1/resumes/00000000-0000-0000-0000-000000000000/export")
    assert response.status_code == 401


def test_export_unknown_resume_returns_404(client: TestClient) -> None:
    headers = _headers(client)
    response = client.get(
        "/api/v1/resumes/00000000-0000-0000-0000-000000000000/export", headers=headers
    )
    assert response.status_code == 404


def test_export_pdf(client: TestClient) -> None:
    headers = _headers(client)
    resume_id = _upload(client, headers)
    response = client.get(f"/api/v1/resumes/{resume_id}/export?format=pdf", headers=headers)
    assert response.status_code == 200
    assert response.headers["content-type"] == PDF_CONTENT_TYPE
    assert response.headers["content-disposition"] == 'attachment; filename="resume.pdf"'
    assert response.content.startswith(b"%PDF")
    assert "Jane Doe" in _pdf_text(response.content)


def test_export_docx(client: TestClient) -> None:
    headers = _headers(client)
    resume_id = _upload(client, headers)
    response = client.get(f"/api/v1/resumes/{resume_id}/export?format=docx", headers=headers)
    assert response.status_code == 200
    assert response.headers["content-type"] == DOCX_CONTENT_TYPE
    assert response.headers["content-disposition"] == 'attachment; filename="resume.docx"'
    assert "Jane Doe" in _docx_text(response.content)


def test_export_defaults_to_pdf(client: TestClient) -> None:
    headers = _headers(client)
    resume_id = _upload(client, headers)
    response = client.get(f"/api/v1/resumes/{resume_id}/export", headers=headers)
    assert response.status_code == 200
    assert response.headers["content-type"] == PDF_CONTENT_TYPE


def test_export_rejects_unknown_format(client: TestClient) -> None:
    headers = _headers(client)
    resume_id = _upload(client, headers)
    response = client.get(f"/api/v1/resumes/{resume_id}/export?format=txt", headers=headers)
    assert response.status_code == 422
